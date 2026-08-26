from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from .issues_models import EquipmentIssue
from core.core_models import Equipment
from django import forms
import pandas as pd  # 用于Excel导入
from datetime import datetime
from .issues_models import EquipmentIssue, PatrolIssue   # ← 新增这一行
from django.contrib import messages   # ← 新增这一行
from notification.models import EmailGroup   # ← 新增这一行
from datetime import date
import json
from django.http import JsonResponse
from .issues_models import PatrolIssue, PatrolCategory
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.shortcuts import render, redirect, get_object_or_404
from .issues_models import EquipmentIssue, EquipmentIssueCategory
from core.core_models import Equipment
from django.contrib import messages
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from .issues_models import EquipmentIssue, EquipmentIssueCategory, ProductionLineConfig



class IssueForm(forms.ModelForm):
    class Meta:
        model = EquipmentIssue
        fields = ['desc', 'occur_date', 'abnormal_work_time', 'work_time_type', 'severity']  # 排除 issue_code


# ==================== 新增：主管下拉列表（以后要加主管只改这里） ====================
SUPERVISOR_CHOICES = [
    "Wind Wang(王海风_PVN)",
    "Xh Li(李興恒_PVN)",
    "Zheng-ping Liu(劉正平_PVN)",
    "Guang-hua Peng(彭光華_PVN)",
]
# ==================== 允许编辑/删除点检记录的用户（白名单） ====================
# 只有这些用户名可以看到并操作“编辑”和“删除”按钮
# 以后想加人或删人，只改这里就行
PATROL_EDIT_DELETE_USERS = [
    'V25020512',
    's20038978',
    'V25018772',
    'perry_chen',
    # 以后可以继续往这里加

]
# 公屏通知权限白名单（可发布/关闭跑马灯公屏信息）
MARQUEE_NOTICE_USERS = [
    'V25020512',
    's20038978',
    'V25018772',
    'perry_chen',
    # 以后要加账号，直接写在这里
]

# 设备异常：可删除的白名单（以后要加账号写这里）
EQUIPMENT_ISSUE_DELETE_USERS = [
    'perry_chen',
    's20038978',
    'V25020512',
    'V25018772',
]
# 兼容旧代码，避免报 NameError
EQUIPMENT_ISSUE_EDIT_USERS = EQUIPMENT_ISSUE_DELETE_USERS



# ==================== 新增：跑马灯日志辅助函数（只在新增记录时使用） ====================
def append_to_marquee_log(request, line):
    """点检记录新增成功后，自动在跑马灯追加一条记录"""
    from django.conf import settings
    import os
    from datetime import datetime

    try:
        # 优先使用用户姓名，没有姓名就用登录账号
        if hasattr(request.user, 'profile') and request.user.profile.name:
            display_name = request.user.profile.name.strip()
        else:
            display_name = request.user.username

        # 构建显示内容
        if line and line.strip():
            log_text = f"{display_name} 在 {line.strip()} 线 点检Điểm kiểm trên dây chuyền"
        else:
            log_text = f"{display_name} 点检记录已上传Ghi chép điểm kiểm đã tải lên"

        # 加上时间戳
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M")
        full_line = f"{timestamp} | {log_text}"

        # 写入 maintenance_log.txt（追加模式）
        log_path = os.path.join(settings.BASE_DIR, 'static', 'maintenance_log.txt')
        os.makedirs(os.path.dirname(log_path), exist_ok=True)

        with open(log_path, 'a', encoding='utf-8') as f:
            f.write(full_line + '\n')

        print(f"[跑马灯] 已成功追加：{full_line}")   # 调试信息
    except Exception as e:
        print(f"[跑马灯] 追加失败：{e}")



@login_required
def issues_list(request, equipment_pk):
    """
    问题列表视图。
    输入：请求和equipment_pk。
    输出：渲染模板。
    功能：显示设备的问题列表，支持过滤。
    交互：通过pk拉取equipment.issues.all()。
    """
    equipment = get_object_or_404(Equipment, pk=equipment_pk)
    issues = equipment.issues.all()
    return render(request, 'issues/issues_list.html', {'issues': issues, 'equipment': equipment})



@login_required
def patrol_list(request):
    """点检问题点记录列表（支持日期 + 主管 + 状态筛选 + 分页）"""

    # ==================== 1. 获取筛选参数 ====================
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    supervisor = request.GET.get('supervisor', '').strip()
    status = request.GET.get('status', '').strip()
    category = request.GET.get('category', '').strip()  # ← 新增类别大饼跳转用

    # ==================== 2. 构建查询 ====================
    issues = PatrolIssue.objects.all().order_by('-date')

    if start_date:
        issues = issues.filter(date__gte=start_date)
    if end_date:
        issues = issues.filter(date__lte=end_date)

    # 新增：主管筛选（支持模糊匹配，兼容各种写法）
    if supervisor:
        issues = issues.filter(supervisor__icontains=supervisor)
    if status:
        issues = issues.filter(status__iexact=status)
    if category:
        # 支持按类别名称筛选（也兼容「未分类」）
        if category == '未分类':
            issues = issues.filter(category__isnull=True)
        else:
            issues = issues.filter(category__name=category)

    # 新增：状态筛选
    if status:
        issues = issues.filter(status__iexact=status)

    # ==================== 3. 分页 ====================
    from django.core.paginator import Paginator
    paginator = Paginator(issues, 20)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    groups = EmailGroup.objects.all().order_by('name')

    return render(request, 'issues/patrol_list.html', {
        'issues': page_obj,
        'groups': groups,
        'today': date.today(),
        'page_obj': page_obj,
        'start_date': start_date,
        'end_date': end_date,
        'selected_supervisor': supervisor,
        'selected_status': status,
        'patrol_edit_users': PATROL_EDIT_DELETE_USERS,  # 保持之前的
        'selected_category': category, # 类别大饼跳用
    })


@login_required
def patrol_create(request):
    """新增点检问题点记录（支持多张照片 + 正确保存到 media 文件夹）"""
    if request.method == 'POST':
        photos = request.FILES.getlist('photos')   # 获取所有上传的文件
        photo_urls = []

        # 获取用户选择的类别ID
        category_id = request.POST.get('category')

        issue = PatrolIssue.objects.create(
            sequence=request.POST.get('sequence'),
            date=request.POST.get('date'),
            line=request.POST.get('line'),
            station=request.POST.get('station'),
            problem_desc=request.POST.get('problem_desc'),
            op_responsible=request.POST.get('op_responsible'),
            supervisor=request.POST.get('supervisor'),
            status=request.POST.get('status'),
            category_id=category_id if category_id else None,  # ← 新增这一行
        )

        # 保存每张照片到磁盘，并记录URL
        for photo in photos:
            if photo:
                # 保存文件
                issue.photo = photo   # 先用单文件字段保存
                issue.save()          # 保存一次
                photo_urls.append(issue.photo.url)   # 记录URL

        issue.photos = photo_urls   # 保存到多照片字段
        issue.save()

        # ==================== 新增：更新跑马灯（只在新增时执行） ====================
        append_to_marquee_log(request, issue.line)

        messages.success(request, f'✅ 点检记录已保存！共上传 {len(photos)} 张照片')
        return redirect('issues:patrol_list')

    # 获取所有启用的类别，按排序显示
    categories = PatrolCategory.objects.filter(is_active=True).order_by('order')

    return render(request, 'issues/patrol_form.html', {
        'supervisor_choices': SUPERVISOR_CHOICES,
        'categories': categories,  # ← 新增
    })


@login_required
def patrol_edit(request, pk):
    """编辑点检问题点记录（支持删除已有照片）—— 增加权限控制"""

    # ==================== 权限检查 ====================
    if request.user.username not in PATROL_EDIT_DELETE_USERS:
        messages.error(request, '❌ 你没有权限编辑点检记录！')
        return redirect('issues:patrol_list')
    # =================================================

    issue = get_object_or_404(PatrolIssue, pk=pk)
    if request.method == 'POST':
        photos_to_delete = request.POST.get('photos_to_delete', '').strip()
        new_photos = request.FILES.getlist('photos')

        # 删除选中的照片
        if photos_to_delete:
            delete_list = [u.strip() for u in photos_to_delete.split(',') if u.strip()]
            issue.photos = [url for url in issue.photos if url not in delete_list]

        # 保存新上传的照片
        for photo in new_photos:
            if photo:
                issue.photo = photo          # 先用单文件字段保存
                issue.save()
                if issue.photo.url not in issue.photos:
                    issue.photos.append(issue.photo.url)

        # 更新其他字段
        issue.date = request.POST.get('date')
        issue.line = request.POST.get('line')
        issue.station = request.POST.get('station')
        issue.problem_desc = request.POST.get('problem_desc')
        issue.op_responsible = request.POST.get('op_responsible')
        issue.supervisor = request.POST.get('supervisor')
        issue.status = request.POST.get('status')
        # 更新类别
        category_id = request.POST.get('category')
        issue.category_id = category_id if category_id else None

        issue.save()

        messages.success(request, f'✅ 编辑成功！已删除 {len(delete_list) if photos_to_delete else 0} 张照片')
        return redirect('issues:patrol_list')

    categories = PatrolCategory.objects.filter(is_active=True).order_by('order')

    return render(request, 'issues/patrol_form.html', {
        'issue': issue,
        'edit': True,
        'supervisor_choices': SUPERVISOR_CHOICES,
        'categories': categories,  # ← 新增
    })



@login_required
def patrol_import(request):
    """点检问题点 Excel 导入（完全匹配你提供的格式）"""
    if request.method == 'POST' and 'excel_file' in request.FILES:
        excel_file = request.FILES['excel_file']
        try:
            df = pd.read_excel(excel_file, sheet_name="问题记录 2026")
            count = 0

            for idx, row in df.iterrows():
                if idx < 1: continue  # 跳过标题行

                date_str = str(row.get('日期', '')).strip()
                if not date_str or date_str in ['日期', 'nan', 'NaN', '']:
                    continue

                try:
                    issue_date = pd.to_datetime(date_str).date()
                except:
                    continue

                PatrolIssue.objects.create(
                    sequence=row.get('顺序'),
                    date=issue_date,
                    line=str(row.get('线别', '')).strip(),
                    station=str(row.get('站别', '')).strip(),
                    problem_desc=str(row.get('问题要改善', '')).strip(),
                    op_responsible=str(row.get('越南负责人（OP）', '')).strip(),
                    supervisor=str(row.get('责任单位主管', '')).strip(),
                    status=str(row.get('情况', 'Open')).strip(),
                    # 照片暂时不支持批量导入（可后续手动编辑上传）
                )
                count += 1

            messages.success(request, f'✅ 成功导入 {count} 条点检问题记录！')
            return redirect('issues:patrol_list')

        except Exception as e:
            messages.error(request, f'❌ 导入失败：{str(e)}')
            return redirect('issues:patrol_list')

    return render(request, 'issues/patrol_import.html', {})


@login_required
def patrol_delete(request, pk):
    """删除一条点检记录 —— 增加权限控制"""

    # ==================== 权限检查 ====================
    if request.user.username not in PATROL_EDIT_DELETE_USERS:
        messages.error(request, '❌ 你没有权限删除点检记录！')
        return redirect('issues:patrol_list')
    # =================================================

    from django.shortcuts import get_object_or_404
    from django.contrib import messages
    issue = get_object_or_404(PatrolIssue, pk=pk)
    issue.delete()
    messages.success(request, f'✅ 记录 {pk} 已删除！')
    return redirect('issues:patrol_list')


@login_required
def global_issues_list(request):
    """
    设备异常分析 - 全局列表
    支持：日期、机种、线体、主管、严重度、关键词 + 异常时间排序 + 分页
    """
    from django.core.paginator import Paginator
    from core.core_models import Equipment

    start_date = request.GET.get('start_date', '').strip()
    end_date = request.GET.get('end_date', '').strip()
    model_type = request.GET.get('model_type', '').strip()
    line = request.GET.get('line', '').strip()
    supervisor = request.GET.get('supervisor', '').strip()
    severity = request.GET.get('severity', '').strip()
    keyword = request.GET.get('keyword', '').strip()
    issue_status = request.GET.get('issue_status', '').strip()
    sort = request.GET.get('sort', '').strip()  # time_desc / time_asc

    issues = EquipmentIssue.objects.select_related(
        'equipment', 'category'
    )

    if start_date:
        issues = issues.filter(occur_date__gte=start_date)
    if end_date:
        issues = issues.filter(occur_date__lte=end_date)
    if model_type:
        issues = issues.filter(equipment__model_type=model_type)
    if line:
        issues = issues.filter(equipment__line=line)
    if supervisor:
        issues = issues.filter(supervisor__icontains=supervisor)
    if severity:
        issues = issues.filter(severity=severity)
    if keyword:
        issues = issues.filter(desc__icontains=keyword)
    if issue_status:
        issues = issues.filter(spare_field_3=issue_status)

    # 排序：默认按日期倒序；可按异常时间
    if sort == 'time_asc':
        issues = issues.order_by('abnormal_work_time', '-occur_date')
    elif sort == 'time_desc':
        issues = issues.order_by('-abnormal_work_time', '-occur_date')
    else:
        issues = issues.order_by('-occur_date', '-id')

    paginator = Paginator(issues, 20)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    model_types = (
        Equipment.objects.exclude(model_type__isnull=True).exclude(model_type='')
        .values_list('model_type', flat=True).distinct().order_by('model_type')
    )
    lines = (
        Equipment.objects.exclude(line__isnull=True).exclude(line='')
        .values_list('line', flat=True).distinct().order_by('line')
    )

    context = {
        'page_obj': page_obj,
        'issues': page_obj,
        'model_types': model_types,
        'lines': lines,
        'supervisor_choices': SUPERVISOR_CHOICES,
        'start_date': start_date,
        'end_date': end_date,
        'selected_model_type': model_type,
        'selected_line': line,
        'selected_supervisor': supervisor,
        'selected_severity': severity,
        'keyword': keyword,
        'issue_status': issue_status,
        'sort': sort,
    }
    return render(request, 'issues/global_issues_list.html', context)


@login_required
def issue_create(request, equipment_pk=None):
    """
    新增设备异常
    机种→类别→线体：来自 ProductionLineConfig
    站别：来自设备总表（按线体）
    """
    # ---------- 配置表数据（给前端联动） ----------
    configs = list(
        ProductionLineConfig.objects.filter(is_active=True)
        .order_by('order', 'model_type', 'category', 'line')
        .values('model_type', 'category', 'line', 'supervisor')
    )

    # 机种只来自配置表（后台加什么，下拉就有什么）
    model_types = []
    for c in configs:
        if c['model_type'] not in model_types:
            model_types.append(c['model_type'])

    # 站别：按线体从设备表汇总 { line: [station, ...] }
    from collections import defaultdict
    # 站别按「机种 + 线别」分组，避免不同机种同名线体混站
    stations_by_line = defaultdict(list)
    for row in Equipment.objects.exclude(line__isnull=True).exclude(line='').exclude(station__isnull=True).exclude(
            station='').values('model_type', 'line', 'station'):
        mt = (row['model_type'] or '').strip()
        line = row['line'].strip()
        st = row['station'].strip()
        key = mt + '||' + line
        if st and st not in stations_by_line[key]:
            stations_by_line[key].append(st)
    for k in stations_by_line:
        stations_by_line[k].sort()

    # 设备列表（选站后匹配 equipment_id）
    equipments = list(
        Equipment.objects.all().values('id', 'code', 'name', 'model_type', 'line', 'station')
    )

    categories = EquipmentIssueCategory.objects.filter(is_active=True).order_by('order')

    if request.method == 'POST':
        model_type = request.POST.get('model_type', '').strip()
        category_line = request.POST.get('line_category', '').strip()  # 产线类别 Auto/manual
        line = request.POST.get('line', '').strip()
        station = request.POST.get('station', '').strip()
        equipment_id = request.POST.get('equipment_id', '').strip()

        desc = request.POST.get('desc', '').strip()
        cause_5w = request.POST.get('spare_field_1', '').strip()  # 异常原因（5个为什么）
        occur_date = request.POST.get('occur_date', '').strip()
        abnormal_work_time = request.POST.get('abnormal_work_time', '0').strip() or '0'
        severity = request.POST.get('severity', '2')
        root_cause = request.POST.get('root_cause', '').strip()
        category_id = request.POST.get('category', '').strip()  # 问题类别（硬件/软件）
        supervisor = request.POST.get('supervisor', '').strip()
        handler = request.POST.get('spare_field_2', '').strip()# 处理人
        issue_status = request.POST.get('spare_field_3', 'Open').strip() or 'Open'# 状态

        if not desc or not occur_date:
            messages.error(request, '异常现象和发生日期为必填！')
            return redirect(request.path)

        if not cause_5w:
            messages.error(request, '请填写异常原因（5个为什么）！')
            return redirect(request.path)

        # ---------- 校验机种 / 线体 ----------
        # 已选中具体设备（从设备详情进来）时，不再强制机种、线体必须在配置表
        if not equipment_id:
            if not model_type or not ProductionLineConfig.objects.filter(
                    is_active=True, model_type=model_type
            ).exists():
                messages.error(request, '机种不在配置表中，请联系管理员在后台添加！')
                return redirect(request.path)

            cfg_ok = ProductionLineConfig.objects.filter(
                is_active=True, model_type=model_type, line=line
            ).exists()
            if not line or not cfg_ok:
                messages.error(request, '请选择配置表中的有效线体！')
                return redirect(request.path)

        # ---------- 匹配设备 ----------
        equipment = None
        if equipment_id:
            equipment = Equipment.objects.filter(id=equipment_id).first()
        if not equipment and line and station:
            equipment = Equipment.objects.filter(line=line, station=station).first()
        if not equipment:
            messages.error(request, '未匹配到设备，请检查线体/站别或设备总表！')
            return redirect(request.path)

        # 无手选主管时，用配置表默认主管
        if not supervisor:
            cfg = ProductionLineConfig.objects.filter(
                is_active=True, model_type=model_type, line=line
            ).first()
            if cfg:
                supervisor = cfg.supervisor or ''

        # 问题码
        today_str = timezone.now().strftime('%Y%m%d')
        last = (
            EquipmentIssue.objects.filter(issue_code__startswith=today_str)
            .order_by('-issue_code').first()
        )
        if last and last.issue_code:
            try:
                seq = int(last.issue_code.split('-')[-1]) + 1
            except Exception:
                seq = 1
        else:
            seq = 1
        issue_code = f"{today_str}-{seq:03d}"
        risk_id = f"RISK-{issue_code}"

        issue = EquipmentIssue.objects.create(
            equipment=equipment,
            issue_code=issue_code,
            desc=desc,
            spare_field_1=cause_5w,  # ← 加上这一行
            occur_date=occur_date,
            severity=int(severity) if str(severity).isdigit() else 2,
            root_cause=root_cause,
            abnormal_work_time=int(abnormal_work_time) if str(abnormal_work_time).isdigit() else 0,
            work_time_type='',
            category_id=category_id if category_id else None,
            risk_id=risk_id,
            supervisor=supervisor,
            spare_field_2=handler,
            spare_field_3=issue_status,
        )

        photos = request.FILES.getlist('photos')
        for f in photos:
            if f and not issue.photo:
                issue.photo = f
                issue.save()
                break

        messages.success(request, f'✅ 已新增设备异常：{issue_code}')
        if equipment_pk:
            return redirect('issues:issues_list', equipment_pk=equipment_pk)
        return redirect('issues:global_issues_list')

    # ---------- 从设备详情进来：预填位置 ----------
    prefill = {}
    if equipment_pk:
        eq = Equipment.objects.filter(pk=equipment_pk).first()
        if eq:
            prefill = {
                'equipment_id': eq.id,
                'model_type': (eq.model_type or '').strip(),
                'line': (eq.line or '').strip(),
                'station': (eq.station or '').strip(),
                'code': eq.code or '',
                'name': eq.name or '',
            }
            # 若设备机种不在配置表下拉里，临时加入，避免选不中
            if prefill['model_type'] and prefill['model_type'] not in model_types:
                model_types.insert(0, prefill['model_type'])

    import json
    context = {
        'model_types': model_types,
        'configs_json': json.dumps(configs, ensure_ascii=False),
        'stations_by_line_json': json.dumps(dict(stations_by_line), ensure_ascii=False),
        'equipments_json': json.dumps(equipments, ensure_ascii=False),
        'categories': categories,
        'supervisor_choices': SUPERVISOR_CHOICES,
        'today': timezone.now().date().isoformat(),
        'prefill': prefill,
        'prefill_json': json.dumps(prefill, ensure_ascii=False),
    }
    return render(request, 'issues/issue_form.html', context)

@login_required
def import_issues(request):
    """
    最终修正版 - 强制读取 'LapTop BB Auto Downtime' sheet
    """
    if request.method == 'POST' and 'excel_file' in request.FILES:
        excel_file = request.FILES['excel_file']
        try:
            # 强制使用正确的 sheet
            sheet_name = 'LapTop BB Auto Downtime'
            df = pd.read_excel(excel_file, sheet_name=sheet_name, header=0)

            print("=== Excel 文件信息 ===")
            print(f"文件名称: {excel_file.name}")
            print(f"强制读取 sheet: {sheet_name}")
            print(f"实际列名列表: {list(df.columns)}")
            print(f"总行数: {len(df)}")

            print("\n=== 开始导入调试（前5行） ===")
            count = 0
            failed = []

            for idx, row in df.iterrows():
                # 打印前5行关键数据
                if idx < 5:
                    print(f"行{idx+2}: 站別='{row.get('站別', '')}' | 設備='{row.get('設備', '')}' | 日期='{row.get('日期', '')}' | 異常現象='{row.get('異常現象及原因', '')}'")

                station = str(row.get('站別', '')).strip()
                equip_name = str(row.get('設備', '')).strip()
                raw_date = str(row.get('日期', '')).strip()
                desc = str(row.get('異常現象及原因', '')).strip()

                if not station or not desc or raw_date in ['日期', 'nan', '', 'NaN']:
                    continue

                # 优先用站別匹配设备
                equipment = Equipment.objects.filter(station__iexact=station).first()

                # 如果站別没找到，再用设备名称模糊匹配
                if not equipment and equip_name:
                    equipment = Equipment.objects.filter(name__icontains=equip_name).first()

                if not equipment:
                    failed.append(f"行{idx+2}: 未找到设备（站別={station}, 名称={equip_name}）")
                    continue

                # 日期解析
                clean_date = raw_date.split('N')[0].split('D')[0].strip()
                try:
                    occur_date = pd.to_datetime(f"2026-{clean_date}", format='%Y-%m/%d').date()
                except:
                    failed.append(f"行{idx+2}: 日期格式错误 ({raw_date})")
                    continue

                time_str = str(row.get('異常時間\n(min)', row.get('異常時間(min)', '0'))).strip()
                minutes = float(time_str) if time_str.replace('.', '', 1).isdigit() else 0

                EquipmentIssue.objects.create(
                    equipment=equipment,
                    issue_code=f"EX-{datetime.now().strftime('%Y%m%d')}-{count+1:03d}",
                    desc=desc,
                    spare_field_1=cause_5w,
                    occur_date=occur_date,
                    abnormal_work_time=int(minutes),
                    work_time_type=str(row.get('异常工时甄别（硬件/软件/其它)', '')),
                    can_import_maintenance=str(row.get('可否导入保养', '')) == 'OK',
                    risk_id=f"RISK-{datetime.now().strftime('%Y%m%d')}-{count+1:03d}",
                    source_excel=excel_file.name
                )
                count += 1

            msg = f'✅ 成功导入 {count} 条问题点！'
            if failed:
                msg += f'\n⚠️ 失败 {len(failed)} 条（见下方）'
                for f in failed[:15]:
                    print(f)
                if len(failed) > 15:
                    print(f"... 还有 {len(failed)-15} 条失败记录")

            return render(request, 'issues/import_form.html', {'success': msg})

        except Exception as e:
            return render(request, 'issues/import_form.html', {'error': f'导入失败：{str(e)}'})

    return render(request, 'issues/import_form.html', {})


@login_required
def patrol_update_problem(request, pk):
    """只更新“问题要改善”字段（双击编辑专用）"""
    if request.method == 'POST':
        issue = get_object_or_404(PatrolIssue, pk=pk)

        # 支持 POST 表单和 JSON 两种方式
        new_desc = request.POST.get('problem_desc')
        if not new_desc and request.body:
            try:
                data = json.loads(request.body)
                new_desc = data.get('problem_desc')
            except:
                pass

        if new_desc is not None:
            issue.problem_desc = new_desc
            issue.save()
            return JsonResponse({'success': True})

    return JsonResponse({'success': False})


@login_required
def patrol_export(request):
    """导出Excel - 只导出勾选的行 + 照片/视频做成可点击超链接"""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    if request.method != 'POST':
        return redirect('issues:patrol_list')

    selected_ids_str = request.POST.get('selected_ids', '')
    if not selected_ids_str:
        return redirect('issues:patrol_list')

    selected_ids = [int(x) for x in selected_ids_str.split(',') if x.strip().isdigit()]
    issues = PatrolIssue.objects.filter(id__in=selected_ids).order_by('-date')

    wb = Workbook()
    ws = wb.active
    ws.title = "点检问题点记录"

    # 列顺序与网页完全一致
    headers = ['日期', '线别', '站别', '问题要改善', '照片/视频', 'OP负责人', '主管', '状态']

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, issue in enumerate(issues, start=2):
        ws.cell(row=row_idx, column=1, value=issue.date.strftime('%Y-%m-%d') if issue.date else '')
        ws.cell(row=row_idx, column=2, value=issue.line)
        ws.cell(row=row_idx, column=3, value=issue.station)
        ws.cell(row=row_idx, column=4, value=issue.problem_desc)

        # 第5列：照片/视频 → 做成超链接
        if issue.photos and len(issue.photos) > 0:
            url = issue.photos[0]  # 取第一张作为链接
            cell = ws.cell(row=row_idx, column=5, value="查看图片/视频")
            cell.hyperlink = url
            cell.style = "Hyperlink"
        elif issue.photo:
            cell = ws.cell(row=row_idx, column=5, value="查看图片/视频")
            cell.hyperlink = issue.photo.url
            cell.style = "Hyperlink"
        else:
            ws.cell(row=row_idx, column=5, value="无")

        ws.cell(row=row_idx, column=6, value=issue.op_responsible)
        ws.cell(row=row_idx, column=7, value=issue.supervisor)
        ws.cell(row=row_idx, column=8, value=issue.status)

        ws.row_dimensions[row_idx].height = 150

    # 列宽调整
    column_widths = [12, 12, 10, 60, 25, 18, 18, 12]
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=点检问题点记录_{len(issues)}条.xlsx'

    wb.save(response)
    return response

# ==================== 公屏通知管理 ====================
@login_required
def marquee_notice_manage(request):
    """
    公屏通知发布/关闭页面
    只有白名单用户可以访问
    """
    from .issues_models import MarqueeNotice

    # 权限检查
    if request.user.username not in MARQUEE_NOTICE_USERS:
        messages.error(request, '你没有权限管理公屏通知！')
        return redirect('issues:patrol_list')

    # 获取当前启用的通知（最多一条）
    current_notice = MarqueeNotice.objects.filter(is_active=True).order_by('-created_at').first()

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'publish':
            content = request.POST.get('content', '').strip()
            if not content:
                messages.error(request, '通知内容不能为空！')
            else:
                # 先把旧的全部关闭
                MarqueeNotice.objects.filter(is_active=True).update(is_active=False)
                # 创建新通知
                MarqueeNotice.objects.create(
                    content=content,
                    is_active=True,
                    created_by=request.user.username
                )
                messages.success(request, f'公屏通知已发布：{content}')
                return redirect('issues:marquee_notice_manage')

        elif action == 'close':
            MarqueeNotice.objects.filter(is_active=True).update(is_active=False)
            messages.success(request, '公屏通知已关闭，跑马灯恢复显示点检日志')
            return redirect('issues:marquee_notice_manage')

    return render(request, 'issues/marquee_notice.html', {
        'current_notice': current_notice,
        'marquee_users': MARQUEE_NOTICE_USERS,
    })



def marquee_content_api(request):
    """
    跑马灯内容接口（方案B：公屏优先，没有才显示点检日志）
    返回纯文本，给 base.html 的 JS 调用
    """
    from .issues_models import MarqueeNotice
    from django.conf import settings
    import os

    # 1. 优先查启用中的公屏通知
    notice = MarqueeNotice.objects.filter(is_active=True).order_by('-created_at').first()
    if notice:
        time_str = notice.created_at.strftime('%m-%d %H:%M')
        msg = f"📢 {notice.content}　|　发布人：{notice.created_by}　|　时间：{time_str}"
        return HttpResponse(msg, content_type='text/plain; charset=utf-8')

    # 2. 没有公屏通知 → 读原来的点检日志
    log_path = os.path.join(settings.BASE_DIR, 'static', 'maintenance_log.txt')
    try:
        if os.path.exists(log_path):
            with open(log_path, 'r', encoding='utf-8') as f:
                lines = [line.strip() for line in f if line.strip()]
            if lines:
                return HttpResponse('🚨 ' + lines[-1], content_type='text/plain; charset=utf-8')
    except Exception:
        pass

    return HttpResponse('🚨 系统运行正常 | 等待最新信息...', content_type='text/plain; charset=utf-8')


# ==================== 设备异常：详情 / 编辑 / 删除 ====================
@login_required
def issue_detail(request, pk):
    """
    设备异常详情（主管检讨页）
    本条信息 + 设备图/异常图 + 近8周记录 + 近8周折线
    """
    from datetime import timedelta
    from django.utils import timezone
    from collections import defaultdict
    import json

    issue = get_object_or_404(
        EquipmentIssue.objects.select_related('equipment', 'category'),
        pk=pk
    )

    # 同一台设备，近8周
    eight_weeks_ago = timezone.now().date() - timedelta(days=56)
    history_list = list(
        EquipmentIssue.objects.filter(
            equipment=issue.equipment,
            occur_date__gte=eight_weeks_ago,
        ).select_related('equipment', 'category').order_by('-occur_date', '-id')
    )

    history_count = len(history_list)
    history_minutes = 0
    history_open = 0
    for row in history_list:
        history_minutes += row.abnormal_work_time or 0
        st = (row.spare_field_3 or 'Open').strip() or 'Open'
        if st != 'Closed':
            history_open += 1

    # 近8个自然周（周一到周日）异常分钟
    today = timezone.now().date()
    this_monday = today - timedelta(days=today.weekday())
    week_starts = []
    i = 7
    while i >= 0:
        week_starts.append(this_monday - timedelta(days=7 * i))
        i -= 1

    week_map = defaultdict(int)
    for row in history_list:
        if not row.occur_date:
            continue
        monday = row.occur_date - timedelta(days=row.occur_date.weekday())
        week_map[monday] += row.abnormal_work_time or 0

    week_labels = []
    week_values = []
    for ws in week_starts:
        week_labels.append('W%02d' % ws.isocalendar()[1])
        week_values.append(week_map.get(ws, 0))

    # 登录就能编辑；删除看白名单
    can_edit = True
    delete_users = []
    try:
        delete_users = EQUIPMENT_ISSUE_DELETE_USERS
    except NameError:
        try:
            delete_users = EQUIPMENT_ISSUE_EDIT_USERS
        except NameError:
            delete_users = []
    can_delete = request.user.username.lower() in [x.lower() for x in delete_users]

    return render(request, 'issues/issue_detail.html', {
        'issue': issue,
        'can_edit': can_edit,
        'can_delete': can_delete,
        'history_list': history_list,
        'history_count': history_count,
        'history_minutes': history_minutes,
        'history_open': history_open,
        'week_labels_json': json.dumps(week_labels, ensure_ascii=False),
        'week_values_json': json.dumps(week_values),
    })


@login_required
def issue_edit(request, pk):
    # """编辑设备异常（白名单）"""
    # if request.user.username not in EQUIPMENT_ISSUE_EDIT_USERS:
    #     messages.error(request, '你没有权限编辑！')
    #     return redirect('issues:global_issues_list')

    issue = get_object_or_404(EquipmentIssue, pk=pk)
    categories = EquipmentIssueCategory.objects.filter(is_active=True).order_by('order')

    if request.method == 'POST':
        issue.desc = request.POST.get('desc', '').strip()
        issue.spare_field_1 = request.POST.get('spare_field_1', '').strip()
        issue.occur_date = request.POST.get('occur_date') or issue.occur_date
        try:
            issue.abnormal_work_time = int(request.POST.get('abnormal_work_time') or 0)
        except Exception:
            issue.abnormal_work_time = 0
        try:
            issue.severity = int(request.POST.get('severity') or 2)
        except Exception:
            issue.severity = 2
        issue.root_cause = request.POST.get('root_cause', '').strip()
        issue.supervisor = request.POST.get('supervisor', '').strip()
        issue.spare_field_2 = request.POST.get('spare_field_2', '').strip()
        issue.spare_field_3 = request.POST.get('spare_field_3', 'Open').strip() or 'Open'
        cat_id = request.POST.get('category', '').strip()
        issue.category_id = cat_id if cat_id else None

        issue.can_import_maintenance = bool(request.POST.get('can_import_maintenance'))
        issue.maintenance_id = request.POST.get('maintenance_id', '').strip()
        issue.can_import_design = bool(request.POST.get('can_import_design'))
        issue.design_id = request.POST.get('design_id', '').strip()
        issue.can_import_training = bool(request.POST.get('can_import_training'))
        issue.training_id = request.POST.get('training_id', '').strip()
        issue.can_import_repair = bool(request.POST.get('can_import_repair'))
        issue.repair_id = request.POST.get('repair_id', '').strip()

        issue.save()
        messages.success(request, '✅ 已保存修改')
        return redirect('issues:issue_detail', pk=issue.pk)

    return render(request, 'issues/issue_edit.html', {
        'issue': issue,
        'categories': categories,
        'supervisor_choices': SUPERVISOR_CHOICES,
    })


@login_required
def issue_delete(request, pk):
    """删除设备异常（白名单，需 POST 确认）"""
    if request.user.username not in EQUIPMENT_ISSUE_EDIT_USERS:
        messages.error(request, '你没有权限删除！')
        return redirect('issues:global_issues_list')

    issue = get_object_or_404(EquipmentIssue, pk=pk)
    if request.method == 'POST':
        code = issue.issue_code
        issue.delete()
        messages.success(request, f'✅ 已删除：{code}')
        return redirect('issues:global_issues_list')

    return render(request, 'issues/issue_confirm_delete.html', {'issue': issue})


@login_required
def equipment_issue_export(request):
    """导出勾选的设备异常记录为 Excel（含异常原因 5 Why）"""
    from django.http import HttpResponse
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, '服务器未安装 openpyxl，请执行: pip install openpyxl')
        return redirect('issues:global_issues_list')

    if request.method != 'POST':
        return redirect('issues:global_issues_list')

    selected_ids_str = request.POST.get('selected_ids', '')
    if not selected_ids_str:
        messages.warning(request, '请先勾选要导出的记录')
        return redirect('issues:global_issues_list')

    selected_ids = [int(x) for x in selected_ids_str.split(',') if x.strip().isdigit()]
    issues = EquipmentIssue.objects.filter(id__in=selected_ids).select_related(
        'equipment', 'category'
    ).order_by('-occur_date')

    wb = Workbook()
    ws = wb.active
    ws.title = "设备异常记录"

    # 表头：异常现象 与 改善方式 之间增加「原因(5Why)」
    headers = [
        '日期', '机种', '线体', '站别', '异常时间(min)',
        '异常现象', '原因(5Why)', '改善方式',
        '类别', '严重度', '主管', '问题码', 'RISK ID'
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for row_idx, issue in enumerate(issues, start=2):
        eq = issue.equipment
        ws.cell(row=row_idx, column=1, value=issue.occur_date.strftime('%Y-%m-%d') if issue.occur_date else '')
        ws.cell(row=row_idx, column=2, value=getattr(eq, 'model_type', '') or '')
        ws.cell(row=row_idx, column=3, value=getattr(eq, 'line', '') or '')
        ws.cell(row=row_idx, column=4, value=getattr(eq, 'station', '') or '')
        ws.cell(row=row_idx, column=5, value=issue.abnormal_work_time or 0)
        ws.cell(row=row_idx, column=6, value=issue.desc or '')
        ws.cell(row=row_idx, column=7, value=issue.spare_field_1 or '')   # 原因 5 Why
        ws.cell(row=row_idx, column=8, value=issue.root_cause or '')      # 改善方式
        ws.cell(row=row_idx, column=9, value=issue.category.name if issue.category else '')
        ws.cell(row=row_idx, column=10, value=issue.get_severity_display())
        ws.cell(row=row_idx, column=11, value=issue.supervisor or '')
        ws.cell(row=row_idx, column=12, value=issue.issue_code or '')
        ws.cell(row=row_idx, column=13, value=issue.risk_id or '')

    # 列宽（13 列）
    for i, width in enumerate([12, 12, 14, 10, 12, 28, 40, 28, 16, 8, 22, 14, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="equipment_issues.xlsx"'
    wb.save(response)
    return response

# ==================== 设备异常分析看板（第二期） ====================
# ==================== 设备异常分析看板（拆分：总览 / 趋势 / 线体） ====================
# ==================== 设备异常分析看板 ====================
def get_analytics_model_types():
    """
    分析看板机种列表：只从产线配置表读取（后台加什么就有什么）
    """
    seen = []
    for row in ProductionLineConfig.objects.filter(is_active=True).order_by('order', 'model_type'):
        if row.model_type not in seen:
            seen.append(row.model_type)
    return seen


@login_required
def equipment_analytics_home(request):
    """
    总览页：
    - 上方：BNC / PRO / Arias 卡片（今日、本周）
    - 下方：三个机种各自的折线图（按天14天 / 按周8周）
    - 点击图上的点 → 直接进线体明细页
    """
    from django.db.models import Sum
    from django.utils import timezone
    from datetime import timedelta
    import json

    today = timezone.localdate()
    week_start = today - timedelta(days=today.weekday())

    # 每个机种可以单独选 day/week，参数：mode_BNC=day&mode_PRO=week ...
    charts = []
    cards = []

    model_types = get_analytics_model_types()
    for mt in model_types:
        # ----- 卡片 -----
        today_qs = EquipmentIssue.objects.filter(
            equipment__model_type__icontains=mt, occur_date=today
        )
        week_qs = EquipmentIssue.objects.filter(
            equipment__model_type__icontains=mt,
            occur_date__gte=week_start, occur_date__lte=today
        )
        cards.append({
            'model_type': mt,
            'today_minutes': today_qs.aggregate(t=Sum('abnormal_work_time'))['t'] or 0,
            'today_count': today_qs.count(),
            'week_minutes': week_qs.aggregate(t=Sum('abnormal_work_time'))['t'] or 0,
            'week_count': week_qs.count(),
        })

        # ----- 折线数据 -----
        mode = request.GET.get(f'mode_{mt}', 'day').strip()
        if mode not in ('day', 'week'):
            mode = 'day'

        labels, values, point_keys = [], [], []

        if mode == 'day':
            for i in range(13, -1, -1):
                d = today - timedelta(days=i)
                total = EquipmentIssue.objects.filter(
                    equipment__model_type__icontains=mt, occur_date=d
                ).aggregate(t=Sum('abnormal_work_time'))['t'] or 0
                labels.append(d.strftime('%m-%d'))
                values.append(total)
                point_keys.append(d.strftime('%Y-%m-%d'))
        else:
            current_monday = week_start
            for i in range(7, -1, -1):
                ws = current_monday - timedelta(weeks=i)
                we = ws + timedelta(days=6)
                total = EquipmentIssue.objects.filter(
                    equipment__model_type__icontains=mt,
                    occur_date__gte=ws, occur_date__lte=we
                ).aggregate(t=Sum('abnormal_work_time'))['t'] or 0
                labels.append(f'W{ws.isocalendar()[1]:02d}')
                values.append(total)
                point_keys.append(ws.strftime('%Y-%m-%d'))

        charts.append({
            'model_type': mt,
            'mode': mode,
            'labels': json.dumps(labels, ensure_ascii=False),
            'values': json.dumps(values),
            'point_keys': json.dumps(point_keys),
        })

    return render(request, 'issues/equipment_analytics_home.html', {
        'cards': cards,
        'charts': charts,
        'today': today,
        'week_start': week_start,
    })

@login_required
def equipment_analytics_trend(request, model):
    """
    趋势页：某个机种的按天14天 / 按周8周 折线图
    """
    from django.db.models import Sum
    from django.utils import timezone
    from datetime import timedelta
    import json

    model = (model or '').strip()
    if model not in get_analytics_model_types():
        messages.error(request, '机种不存在')
        return redirect('issues:equipment_analytics')

    mode = request.GET.get('mode', 'day').strip()
    if mode not in ('day', 'week'):
        mode = 'day'

    today = timezone.localdate()
    week_start = today - timedelta(days=today.weekday())

    labels, values = [], []
    # 同时保留「完整日期 / 周起止」，方便点击跳转
    point_keys = []

    if mode == 'day':
        for i in range(13, -1, -1):
            d = today - timedelta(days=i)
            total = EquipmentIssue.objects.filter(
                equipment__model_type__icontains=model, occur_date=d
            ).aggregate(t=Sum('abnormal_work_time'))['t'] or 0
            labels.append(d.strftime('%m-%d'))
            values.append(total)
            point_keys.append(d.strftime('%Y-%m-%d'))  # 跳转用完整日期
    else:
        current_monday = week_start
        for i in range(7, -1, -1):
            ws = current_monday - timedelta(weeks=i)
            we = ws + timedelta(days=6)
            total = EquipmentIssue.objects.filter(
                equipment__model_type__icontains=model,
                occur_date__gte=ws, occur_date__lte=we
            ).aggregate(t=Sum('abnormal_work_time'))['t'] or 0
            labels.append(f'W{ws.isocalendar()[1]:02d}')
            values.append(total)
            point_keys.append(ws.strftime('%Y-%m-%d'))  # 用该周周一作为 point

    return render(request, 'issues/equipment_analytics_trend.html', {
        'model': model,
        'mode': mode,
        'today': today,
        'chart_labels': json.dumps(labels, ensure_ascii=False),
        'chart_values': json.dumps(values),
        'point_keys': json.dumps(point_keys),
    })


@login_required
def equipment_analytics_lines(request, model, mode, point):
    """
    机种某天/某周四宫格：线体柱图 + 类别环 + 站别Pareto + 状态
    """
    from datetime import timedelta, datetime
    from collections import defaultdict
    import json

    model = (model or '').strip()
    mode = (mode or '').strip()
    point = (point or '').strip()

    if model not in get_analytics_model_types() or mode not in ('day', 'week'):
        messages.error(request, '参数错误')
        return redirect('issues:equipment_analytics')

    try:
        base = datetime.strptime(point, '%Y-%m-%d').date()
    except Exception:
        messages.error(request, '日期参数错误')
        return redirect('issues:equipment_analytics')

    if mode == 'day':
        date_start = date_end = base
        title_point = base.strftime('%Y-%m-%d')
    else:
        date_start = base
        date_end = base + timedelta(days=6)
        title_point = 'W%02d (%s ~ %s)' % (base.isocalendar()[1], date_start, date_end)

    qs = EquipmentIssue.objects.filter(
        equipment__model_type__icontains=model,
        occur_date__gte=date_start,
        occur_date__lte=date_end,
    ).select_related('equipment', 'category')

    line_bucket = defaultdict(lambda: {'minutes': 0, 'count': 0})
    cat_bucket = defaultdict(lambda: {'minutes': 0, 'count': 0})
    st_bucket = defaultdict(lambda: {'minutes': 0, 'count': 0})
    status_bucket = defaultdict(lambda: {'minutes': 0, 'count': 0})
    total_minutes = 0
    total_count = 0

    for issue in qs:
        mins = issue.abnormal_work_time or 0
        total_minutes += mins
        total_count += 1

        line_name = (issue.equipment.line or '未填线体').strip() or '未填线体'
        line_bucket[line_name]['minutes'] += mins
        line_bucket[line_name]['count'] += 1

        cat_name = issue.category.name if issue.category else '未分类Chưa phân loại'
        cat_bucket[cat_name]['minutes'] += mins
        cat_bucket[cat_name]['count'] += 1

        st = (issue.equipment.station or '未填站别').strip() or '未填站别'
        st_bucket[st]['minutes'] += mins
        st_bucket[st]['count'] += 1

        stt = (issue.spare_field_3 or 'Open').strip() or 'Open'
        status_bucket[stt]['minutes'] += mins
        status_bucket[stt]['count'] += 1

    line_details = [
        {'line': k, 'minutes': v['minutes'], 'count': v['count']}
        for k, v in sorted(line_bucket.items(), key=lambda x: -x[1]['minutes'])
    ]
    category_details = [
        {'name': k, 'minutes': v['minutes'], 'count': v['count']}
        for k, v in sorted(cat_bucket.items(), key=lambda x: -x[1]['minutes'])
    ]
    station_details = [
        {'station': k, 'minutes': v['minutes'], 'count': v['count']}
        for k, v in sorted(st_bucket.items(), key=lambda x: -x[1]['minutes'])[:10]
    ]
    status_order = ['Open', 'Ongoing', 'Closed']
    status_details = []
    for name in status_order:
        if name in status_bucket:
            status_details.append({
                'name': name,
                'minutes': status_bucket[name]['minutes'],
                'count': status_bucket[name]['count'],
            })
    for name, v in status_bucket.items():
        if name not in status_order:
            status_details.append({'name': name, 'minutes': v['minutes'], 'count': v['count']})

    return render(request, 'issues/equipment_analytics_lines.html', {
        'model': model,
        'mode': mode,
        'point': point,
        'title_point': title_point,
        'total_minutes': total_minutes,
        'total_count': total_count,
        'line_details': line_details,
        'line_labels_json': json.dumps([x['line'] for x in line_details], ensure_ascii=False),
        'line_values_json': json.dumps([x['minutes'] for x in line_details]),
        'cat_labels_json': json.dumps([x['name'] for x in category_details], ensure_ascii=False),
        'cat_values_json': json.dumps([x['minutes'] for x in category_details]),
        'st_labels_json': json.dumps([x['station'] for x in station_details], ensure_ascii=False),
        'st_values_json': json.dumps([x['minutes'] for x in station_details]),
        'status_labels_json': json.dumps([x['name'] for x in status_details], ensure_ascii=False),
        'status_values_json': json.dumps([x['minutes'] for x in status_details]),
        'status_details': status_details,
        'station_details': station_details,
    })


@login_required
def equipment_analytics_stations(request, model, mode, point, line):
    """
    某机种 + 某天/周 + 某线体 四宫格：站别 + 类别 + 状态 + 记录
    """
    from datetime import datetime, timedelta
    from collections import defaultdict
    import json

    model = (model or '').strip()
    mode = (mode or '').strip()
    point = (point or '').strip()
    line = (line or '').strip()

    if model not in get_analytics_model_types() or mode not in ('day', 'week'):
        messages.error(request, '参数错误')
        return redirect('issues:equipment_analytics')

    try:
        base = datetime.strptime(point, '%Y-%m-%d').date()
    except Exception:
        messages.error(request, '日期参数错误')
        return redirect('issues:equipment_analytics_lines', model=model, mode=mode, point=point)

    if mode == 'day':
        date_start = date_end = base
        title_point = base.strftime('%Y-%m-%d')
    else:
        date_start = base
        date_end = base + timedelta(days=6)
        title_point = 'W%02d (%s ~ %s)' % (base.isocalendar()[1], date_start, date_end)

    qs = EquipmentIssue.objects.filter(
        equipment__model_type__icontains=model,
        equipment__line=line,
        occur_date__gte=date_start,
        occur_date__lte=date_end,
    ).select_related('equipment', 'category').order_by('-abnormal_work_time', '-occur_date')

    station_bucket = defaultdict(lambda: {'minutes': 0, 'count': 0})
    category_bucket = defaultdict(lambda: {'minutes': 0, 'count': 0})
    status_bucket = defaultdict(lambda: {'minutes': 0, 'count': 0})
    total_minutes = 0
    total_count = 0

    for issue in qs:
        mins = issue.abnormal_work_time or 0
        total_minutes += mins
        total_count += 1

        st = (issue.equipment.station or '未填站别').strip() or '未填站别'
        station_bucket[st]['minutes'] += mins
        station_bucket[st]['count'] += 1

        cat_name = issue.category.name if issue.category else '未分类Chưa phân loại'
        category_bucket[cat_name]['minutes'] += mins
        category_bucket[cat_name]['count'] += 1

        stt = (issue.spare_field_3 or 'Open').strip() or 'Open'
        status_bucket[stt]['minutes'] += mins
        status_bucket[stt]['count'] += 1

    station_details = [
        {'station': k, 'minutes': v['minutes'], 'count': v['count']}
        for k, v in sorted(station_bucket.items(), key=lambda x: -x[1]['minutes'])
    ]
    category_details = [
        {'name': k, 'minutes': v['minutes'], 'count': v['count']}
        for k, v in sorted(category_bucket.items(), key=lambda x: -x[1]['minutes'])
    ]
    status_order = ['Open', 'Ongoing', 'Closed']
    status_details = []
    for name in status_order:
        if name in status_bucket:
            status_details.append({
                'name': name,
                'minutes': status_bucket[name]['minutes'],
                'count': status_bucket[name]['count'],
            })
    for name, v in status_bucket.items():
        if name not in status_order:
            status_details.append({'name': name, 'minutes': v['minutes'], 'count': v['count']})

    record_list = list(qs[:20])

    return render(request, 'issues/equipment_analytics_stations.html', {
        'model': model,
        'mode': mode,
        'point': point,
        'line': line,
        'title_point': title_point,
        'total_minutes': total_minutes,
        'total_count': total_count,
        'station_details': station_details,
        'category_details': category_details,
        'status_details': status_details,
        'record_list': record_list,
        'st_labels_json': json.dumps([x['station'] for x in station_details], ensure_ascii=False),
        'st_values_json': json.dumps([x['minutes'] for x in station_details]),
        'cat_labels_json': json.dumps([x['name'] for x in category_details], ensure_ascii=False),
        'cat_values_json': json.dumps([x['minutes'] for x in category_details]),
        'status_labels_json': json.dumps([x['name'] for x in status_details], ensure_ascii=False),
        'status_values_json': json.dumps([x['minutes'] for x in status_details]),
    })